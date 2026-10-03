"""
Comprueba que compras_excel.py calcula EXACTAMENTE lo mismo que el Excel GESTION_CORRECTA.

    python verificar_contra_excel.py RUTA/GESTION_CORRECTA.xlsx

Lee los datos de entrada del Excel (códigos, kilos, precios, fechas) y los resultados que el
propio Excel tiene guardados en cada celda con fórmula, recalcula todo con compras_excel.py y
compara celda a celda. Al final lista también los datos raros encontrados en la hoja COMPRAS.

El Excel tiene datos reales de Marinafisk: NO se sube al repositorio (ver .gitignore).
Solo lee el archivo; no lo modifica.
"""

import sys
from collections import defaultdict
from datetime import datetime

import openpyxl

import compras_excel as cx

TOL = 1e-9


def igual(a, b):
    if a is None:
        a = cx.VACIO
    if b is None:
        b = cx.VACIO
    if isinstance(a, datetime):
        a = a.date()
    if isinstance(b, datetime):
        b = b.date()
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= TOL * max(1.0, abs(a), abs(b))
    return a == b


class Contador:
    def __init__(self):
        self.ok = defaultdict(int)
        self.mal = defaultdict(list)

    def comprobar(self, nombre, celda, esperado, obtenido):
        if igual(esperado, obtenido):
            self.ok[nombre] += 1
        else:
            self.mal[nombre].append((celda, esperado, obtenido))

    def informe(self):
        total_mal = 0
        print("\n=== COMPARACIÓN CON EL EXCEL ===")
        for nombre in sorted(set(self.ok) | set(self.mal)):
            n_mal = len(self.mal[nombre])
            total_mal += n_mal
            estado = "OK " if n_mal == 0 else "MAL"
            print(f"[{estado}] {nombre}: {self.ok[nombre]} iguales, {n_mal} distintas")
            for celda, e, o in self.mal[nombre][:5]:
                print(f"        {celda}: Excel={e!r}  referencia={o!r}")
        return total_mal


def main(ruta):
    vals = openpyxl.load_workbook(ruta, data_only=True)
    forms = openpyxl.load_workbook(ruta)
    c = Contador()

    # Un archivo guardado por un programa que no es Excel (p. ej. openpyxl) puede no llevar los
    # resultados de las fórmulas: Excel los recalcula al abrirlo, pero aquí no hay con qué comparar.
    muestra = [(r, col) for r in range(3, 200) for col in (5, 12, 15)
               if isinstance(forms["COMPRAS"].cell(r, col).value, str)
               and forms["COMPRAS"].cell(r, col).value.startswith("=")]
    if muestra and all(vals["COMPRAS"].cell(r, col).value is None for r, col in muestra):
        print("ESTE ARCHIVO NO TIENE GUARDADOS LOS RESULTADOS DE LAS FÓRMULAS.\n"
              "Ábrelo en Excel, pulsa Guardar (así Excel los calcula y los guarda) y vuelve a ejecutar.\n"
              "No se compara nada para no dar diferencias falsas.")
        return -1

    # --- Catálogos -----------------------------------------------------------
    hp = vals["PROVEEDORES"]
    proveedores = {}
    for r in range(3, hp.max_row + 1):
        cod = hp.cell(r, 1).value
        if cod is not None:
            proveedores[cod] = {"nombre": hp.cell(r, 2).value, "op2": hp.cell(r, 3).value}

    hpr = vals["PRODUCTOS"]
    productos = {}
    for r in range(3, hpr.max_row + 1):
        cod = hpr.cell(r, 1).value
        if cod is not None:
            productos[cod] = {"descripcion": hpr.cell(r, 2).value, "fila": r}

    # --- COMPRAS: cada línea ---------------------------------------------------
    hc = vals["COMPRAS"]
    hcf = forms["COMPRAS"]
    ultima = max(r for r in range(3, hc.max_row + 1)
                 if isinstance(hcf.cell(r, 12).value, str) and hcf.cell(r, 12).value.startswith("=IF("))
    compras = []
    for r in range(3, ultima + 1):
        g = lambda col: hc.cell(r, col).value
        cod_prov = g(4)
        if isinstance(hcf.cell(r, 4).value, datetime):  # número con formato de fecha (ver anomalías)
            cod_prov = int(round((hcf.cell(r, 4).value - datetime(1899, 12, 30)).days))
        res = cx.calcular_linea(cod_prov, g(7), g(10), g(11), proveedores, productos)
        for col, attr in [(5, "proveedor"), (6, "op2_flag"), (8, "descripcion"), (12, "base_zgz"),
                          (13, "base_zgz_iva"), (14, "op2"), (15, "base_real"), (16, "iva"), (17, "total_fact")]:
            c.comprobar(f"COMPRAS col {openpyxl.utils.get_column_letter(col)} ({hc.cell(2, col).value})",
                        f"{openpyxl.utils.get_column_letter(col)}{r}", g(col), getattr(res, attr))
        compras.append({"fila": r, "partida": g(1), "fecha": g(2), "alb": g(3), "cod_prov": cod_prov,
                        "cod_prod": g(7), "cajas": g(9), "kilos": g(10), "eur_kg": g(11),
                        "base_real": res.base_real, "T": g(20), "V": g(22), "W": g(23), "X": g(24),
                        "S": g(19), "U": g(21)})

    # --- COMPRAS: partida sugerida (columna W) --------------------------------
    hoy = None
    for x in compras:
        if x["T"] not in (None, ""):
            # HOY() del Excel = día en que se recalculó por última vez; solo importa si B está vacía
            hoy = hoy or (x["fecha"].date() if isinstance(x["fecha"], datetime) else None)
            obt = cx.partida_sugerida_excel(compras, x["fila"], x["T"], x["fecha"], x["X"], hoy)
            c.comprobar("COMPRAS col W (partida sugerida, fórmula tal cual)", f"W{x['fila']}", x["W"], obt)

    # --- PRODUCTOS: última fecha, coste, PVP1, PVP2 ----------------------------
    for cod, p in productos.items():
        r = p["fila"]
        m = cx.ultima_fecha_compra(cod, compras)
        h = cx.coste_producto(cod, compras)
        c.comprobar("PRODUCTOS M (última fecha compra)", f"M{r}", hpr.cell(r, 13).value, m)
        c.comprobar("PRODUCTOS H (coste)", f"H{r}", hpr.cell(r, 8).value, h)
        c.comprobar("PRODUCTOS I (PVP1)", f"I{r}", hpr.cell(r, 9).value,
                    cx.pvp(h, cx.SUMA_PVP1, hpr.cell(r, 24).value))
        c.comprobar("PRODUCTOS J (PVP2)", f"J{r}", hpr.cell(r, 10).value,
                    cx.pvp(h, cx.SUMA_PVP2, hpr.cell(r, 25).value))

    # --- PRECIO MEDIO ---------------------------------------------------------
    hm = vals["PRECIO MEDIO"]
    desde, hasta = hm["C4"].value, hm["D4"].value
    for r in range(11, hm.max_row + 1):
        cod = hm.cell(r, 1).value
        if cod is None or hm.cell(r, 3).value is None:
            continue
        pm = cx.precio_medio(cod, compras, desde, hasta, hm.cell(r, 8).value)
        c.comprobar("PRECIO MEDIO C (total kg)", f"C{r}", hm.cell(r, 3).value, pm["kg"])
        c.comprobar("PRECIO MEDIO D (total base real)", f"D{r}", hm.cell(r, 4).value, pm["base_real"])
        c.comprobar("PRECIO MEDIO E (precio medio)", f"E{r}", hm.cell(r, 5).value, pm["precio_medio"])
        c.comprobar("PRECIO MEDIO F (nº compras)", f"F{r}", hm.cell(r, 6).value, pm["lineas"])

    # --- PANEL COMPRAS: totales por proveedor ----------------------------------
    hpa = vals["PANEL COMPRAS"]
    d1, d2 = hpa["C4"].value, hpa["G4"].value
    for r in range(2, hpa.max_row + 1):
        cod = hpa.cell(r, 14).value
        if cod in (None, "", 0):
            continue
        kg, base = cx.totales_proveedor(cod, compras, d1, d2)
        c.comprobar("PANEL COMPRAS P (kilos proveedor)", f"P{r}", hpa.cell(r, 16).value, kg)
        c.comprobar("PANEL COMPRAS Q (importe proveedor)", f"Q{r}", hpa.cell(r, 17).value, base)

    total_mal = c.informe()
    formulas_que_faltan(forms, ultima)
    anomalias(compras, hcf, proveedores, productos)
    return total_mal


def formulas_que_faltan(forms, ultima_compra):
    """Celdas que deberían tener fórmula y tienen un valor fijo o están vacías (punto 16)."""
    print("\n=== FÓRMULAS QUE FALTAN (valor fijo o vacío donde debería haber fórmula) ===")
    revisar = [
        ("COMPRAS", "EFHLMNOPQ", 3, ultima_compra, lambda h, r: True),
        ("PRODUCTOS", "HIJMN", 3, forms["PRODUCTOS"].max_row, lambda h, r: h.cell(r, 1).value is not None),
        ("PRECIO MEDIO", "CDEF", 11, forms["PRECIO MEDIO"].max_row, lambda h, r: h.cell(r, 1).value is not None),
    ]
    for hoja, cols, desde, hasta, aplica in revisar:
        h = forms[hoja]
        faltan = []
        for r in range(desde, hasta + 1):
            if not aplica(h, r):
                continue
            for col in cols:
                x = h[f"{col}{r}"].value
                if not (isinstance(x, str) and x.startswith("=")):
                    faltan.append((f"{col}{r}", x, h.cell(r, 1).value if hoja != "COMPRAS" else h.cell(r, 7).value))
        print(f"- {hoja}: {len(faltan)}")
        for f in faltan[:12]:
            print(f"      {f[0]} = {f[1]!r}   (código {f[2]})")


def anomalias(compras, hcf, proveedores, productos):
    """Datos raros en COMPRAS que conviene revisar (no son diferencias de cálculo)."""
    print("\n=== DATOS A REVISAR EN LA HOJA COMPRAS ===")
    con_datos = [x for x in compras if any(x[k] not in (None, "") for k in ("cod_prov", "cod_prod", "kilos", "eur_kg"))]

    def lista(titulo, filas):
        print(f"- {titulo}: {len(filas)}")
        for f in filas[:12]:
            print(f"      {f}")

    lista("Líneas con kilos/precio pero SIN número de partida",
          [(x["fila"], x["fecha"].strftime("%d/%m/%Y") if x["fecha"] else None, x["cod_prov"], x["cod_prod"], x["kilos"], x["eur_kg"])
           for x in con_datos if x["partida"] in (None, "")])
    lista("Líneas SIN proveedor (no se le aplica el 2 % OP aunque sea de lonja)",
          [(x["fila"], x["cod_prod"], x["kilos"], x["eur_kg"]) for x in con_datos if x["cod_prov"] in (None, "")])
    lista("Líneas SIN producto",
          [(x["fila"], x["cod_prov"], x["kilos"], x["eur_kg"]) for x in con_datos if x["cod_prod"] in (None, "")])
    lista("Código de proveedor que no existe en PROVEEDORES",
          [(x["fila"], x["cod_prov"]) for x in con_datos if x["cod_prov"] not in (None, "") and x["cod_prov"] not in proveedores])
    lista("Código de producto que no existe en PRODUCTOS",
          [(x["fila"], x["cod_prod"]) for x in con_datos if x["cod_prod"] not in (None, "") and cx._buscar_producto(x["cod_prod"], productos) is None])
    lista("Código escrito a mano DISTINTO del que muestra el buscador de esa misma fila",
          [(x["fila"], f"prov {x['cod_prov']} vs buscador {x['T']}", f"prod {x['cod_prod']} vs buscador {x['V']}")
           for x in con_datos
           if (x["T"] not in (None, "") and x["T"] != x["cod_prov"]) or (x["V"] not in (None, "") and x["V"] != x["cod_prod"])])
    lista("Número de partida con decimales",
          [(x["fila"], x["partida"]) for x in compras if isinstance(x["partida"], float) and not x["partida"].is_integer()])
    lista("Columna CAJAS con una suma de pesos dentro",
          [(x["fila"], hcf.cell(x["fila"], 9).value, x["kilos"]) for x in compras
           if isinstance(hcf.cell(x["fila"], 9).value, str) and hcf.cell(x["fila"], 9).value.startswith("=")])
    lista("Cajas con decimales",
          [(x["fila"], x["cajas"]) for x in compras if isinstance(x["cajas"], float) and not x["cajas"].is_integer()])
    lista("Código de proveedor o albarán guardado con formato de FECHA",
          [(x["fila"], openpyxl.utils.get_column_letter(col), hcf.cell(x["fila"], col).value)
           for x in compras for col in (3, 4) if isinstance(hcf.cell(x["fila"], col).value, datetime)])
    lista("Precio 0 €/kg (cuenta en el coste medio del producto)",
          [(x["fila"], x["cod_prod"], x["kilos"]) for x in con_datos if x["eur_kg"] == 0])

    grupos = defaultdict(set)
    for x in compras:
        if x["partida"] not in (None, ""):
            grupos[x["partida"]].add((x["cod_prov"], x["fecha"].date() if x["fecha"] else None))
    lista("Partidas con más de un proveedor o fecha",
          [(p, sorted(map(str, s))) for p, s in grupos.items() if len(s) > 1])

    print("\nNota: la columna B (FECHA) tiene formato", repr(hcf.cell(3, 2).number_format),
          "(mes/día/año) en este archivo.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(1 if main(sys.argv[1]) else 0)
